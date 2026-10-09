"""Build the complete Homework 1 solution report as a polished DOCX.

The report follows the `standard_business_brief` document preset.  It is later
converted to PDF with Microsoft Word and visually inspected page by page.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "deliverables"
OUT_DIR.mkdir(exist_ok=True)
DOCX_PATH = OUT_DIR / "homework-1-solutions.docx"
METRICS_PATH = ROOT / "task5_pinn" / "results" / "metrics.json"
PLOT_PATH = ROOT / "task5_pinn" / "results" / "best_test_rollout.png"
REPO_URL = "https://github.com/Maks-Grigoryan/photonicsai-homework-1"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "20252B"
MUTED = "5E6872"
LIGHT_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
BORDER = "CBD5E1"
WHITE = "FFFFFF"


def set_cell_text(cell, text: str, *, bold: bool = False, align="left", size=9.4):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    run = paragraph.add_run(str(text))
    set_run(run, size=size, bold=bold)


def set_run(run, *, size=11, bold=None, italic=None, color=INK, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120):
    assert sum(widths_dxa) == 9360, widths_dxa
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    set_table_borders(table)


def add_table(doc, headers: list[str], rows: list[list[object]], widths: list[int], aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.allow_autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, align="center", size=9.2)
        set_cell_shading(table.rows[0].cells[index], LIGHT_FILL)
    set_repeat_table_header(table.rows[0])
    if aligns is None:
        aligns = ["left"] * len(headers)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_text(cells[index], str(value), align=aligns[index], size=9.2)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run(run, size=9, color=MUTED)


def add_bottom_border(paragraph, color=BLUE, size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_text(doc, text: str, *, bold_prefix: str | None = None, style=None):
    paragraph = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        set_run(paragraph.add_run(bold_prefix), bold=True)
        set_run(paragraph.add_run(text[len(bold_prefix) :]))
    else:
        set_run(paragraph.add_run(text))
    return paragraph


def add_equation(doc, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run(run, size=11.3, italic=True, font="Cambria Math", color=DARK_BLUE)
    return paragraph


def add_code(doc, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F7F9FB")
    p_pr.append(shd)
    run = paragraph.add_run(text)
    set_run(run, size=8.4, font="Consolas", color="243447")
    return paragraph


def add_numbered(doc, text: str, num_id: int):
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    set_run(paragraph.add_run(text))
    return paragraph


def create_numbering(doc, fmt: str, text: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend((tabs, ind, spacing))
    level.extend((start, num_fmt, lvl_text, lvl_jc, p_pr))
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def heading(doc, text: str, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def page_break(doc):
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        1: (16, BLUE, 16, 8),
        2: (13, BLUE, 12, 6),
        3: (12, DARK_BLUE, 8, 4),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Quiet running header and right-aligned page number.
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run(hp.add_run("PhotonicsAI  |  Homework 1"), size=8.5, color=MUTED, bold=True)
    set_run(hp.add_run("\tComplete solutions"), size=8.5, color=MUTED)
    tab_stops = hp.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(fp.add_run("Page "), size=9, color=MUTED)
    add_page_number(fp)

    doc.core_properties.title = "PhotonicsAI Homework 1 — Complete Solutions"
    doc.core_properties.subject = "AI, Neural Networks, LLMs, and Agentic AI Systems"
    doc.core_properties.author = "Homework 1 submission"
    doc.core_properties.keywords = "neural networks, softmax, attention, PINN, LangGraph"


def cover_page(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(32)
    kicker = doc.add_paragraph()
    set_run(kicker.add_run("PHOTONICSAI  |  AI-DRIVEN SCIENCE"), size=10.5, bold=True, color=BLUE)
    kicker.paragraph_format.space_after = Pt(7)
    title = doc.add_paragraph()
    set_run(title.add_run("Homework 1"), size=28, bold=True, color="17324D")
    title.paragraph_format.space_after = Pt(4)
    subtitle = doc.add_paragraph()
    set_run(
        subtitle.add_run("Complete worked solutions and reproducible experiments"),
        size=14,
        color=DARK_BLUE,
    )
    subtitle.paragraph_format.space_after = Pt(20)

    for label, value in [
        ("Topics", "Neural networks, parameter counting, softmax, attention, PINNs, agentic systems"),
        ("Due date", "11 October 2026"),
        ("Dataset", "Silverbox / SNLS80mV"),
        ("Code", REPO_URL),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        set_run(p.add_run(f"{label}: "), bold=True, color="17324D")
        if label == "Code":
            add_hyperlink(p, value, value)
        else:
            set_run(p.add_run(value), color=INK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(20)
    add_bottom_border(rule)
    lead = doc.add_paragraph()
    lead.paragraph_format.space_before = Pt(8)
    lead.paragraph_format.space_after = Pt(10)
    set_run(
        lead.add_run(
            "This report gives every requested weight, bias, shape, proof, free-run metric, "
            "and agent trace. Numerical claims in Tasks 5–6 come from the included scripts, "
            "not from hand-entered estimates."
        ),
        size=11.5,
        color="334155",
        italic=True,
    )
    add_text(
        doc,
        "Reproducibility note. The official dataset is not committed to the repository. Place "
        "SNLS80mV.csv under data/SilverboxFiles/ and run the commands in the final section.",
        bold_prefix="Reproducibility note.",
    )
    page_break(doc)


def task1(doc):
    heading(doc, "Task 1. Two networks a line cannot be", 1)
    add_text(
        doc,
        "Both targets are piecewise linear on the requested domain, but neither is a single "
        "affine function. ReLU is useful because each hidden unit contributes one hinge. The "
        "constructions below use two hidden units, and two is minimal in both parts.",
    )

    heading(doc, "(a) Absolute value", 2)
    add_text(
        doc,
        "Use one hinge for the positive half-line and a second hinge for the negative half-line:",
    )
    add_equation(doc, "h₁ = ReLU(x),     h₂ = ReLU(−x),     ŷ = h₁ + h₂.")
    add_text(
        doc,
        "In matrix form, the hidden layer and output layer are",
    )
    add_equation(
        doc,
        "W⁽¹⁾ = [1; −1],   b⁽¹⁾ = [0; 0],   σ⁽¹⁾ = ReLU;      W⁽²⁾ = [1  1],   b⁽²⁾ = 0,   σ⁽²⁾(z)=z.",
    )
    add_text(
        doc,
        "If x ≥ 0, then h₁=x and h₂=0. If x<0, then h₁=0 and h₂=−x. Therefore ŷ=|x| for every real x.",
    )
    add_table(
        doc,
        ["x", "z₁=x", "h₁", "z₂=−x", "h₂", "ŷ"],
        [
            ["−2", "−2", "0", "2", "2", "2"],
            ["−0.5", "−0.5", "0", "0.5", "0.5", "0.5"],
            ["0", "0", "0", "0", "0", "0"],
            ["0.5", "0.5", "0.5", "−0.5", "0", "0.5"],
            ["3", "3", "3", "−3", "0", "3"],
        ],
        [1200, 1500, 1500, 1500, 1500, 2160],
        ["center"] * 6,
    )
    add_text(
        doc,
        "Requested negative-input trace. At x=−2, unit 1 has pre-activation z₁=1·(−2)+0=−2 "
        "and activation h₁=ReLU(−2)=0. Unit 2 has z₂=(−1)·(−2)+0=2 and h₂=ReLU(2)=2. "
        "The identity output is 1·0+1·2+0=2.",
        bold_prefix="Requested negative-input trace.",
    )
    add_text(
        doc,
        "Minimality. With one hidden ReLU, a linear output has the form A·ReLU(wx+b)+c. Its "
        "slope is zero on one side of its hinge and Aw on the other. The function |x| instead "
        "has two nonzero slopes, −1 and +1, so one hidden unit cannot work. Thus the minimum is two.",
        bold_prefix="Minimality.",
    )

    page_break(doc)
    heading(doc, "(b) Exclusive OR", 2)
    add_text(doc, "On the Boolean square, XOR is the absolute difference |x₁−x₂|. Reuse the two-hinge construction:")
    add_equation(doc, "h₁ = ReLU(x₁−x₂),     h₂ = ReLU(x₂−x₁),     ŷ = h₁+h₂.")
    add_equation(
        doc,
        "W⁽¹⁾ = [[1,−1],[−1,1]],   b⁽¹⁾=[0,0],   σ⁽¹⁾=ReLU;      W⁽²⁾=[1,1],   b⁽²⁾=0,   σ⁽²⁾(z)=z.",
    )
    add_table(
        doc,
        ["x₁", "x₂", "z₁", "h₁", "z₂", "h₂", "ŷ"],
        [
            [0, 0, 0, 0, 0, 0, 0],
            [0, 1, "−1", 0, 1, 1, 1],
            [1, 0, 1, 1, "−1", 0, 1],
            [1, 1, 0, 0, 0, 0, 0],
        ],
        [900, 900, 1200, 1200, 1200, 1200, 2760],
        ["center"] * 7,
    )
    add_text(
        doc,
        "Why one neuron is insufficient. The two positive XOR points, (1,0) and (0,1), are on "
        "opposite corners from the two negative points, (0,0) and (1,1); no single affine "
        "decision boundary can separate those two diagonal pairs. The two-ReLU network above "
        "therefore uses the minimum number of hidden units.",
        bold_prefix="Why one neuron is insufficient.",
    )


def task2(doc):
    page_break(doc)
    heading(doc, "Task 2. Parameter count and layer shapes", 1)
    add_text(
        doc,
        "For every convolution or pool I apply Hout=⌊(H+2p−k)/s⌋+1. A convolution contributes "
        "k²CinCout+Cout parameters; a dense layer contributes NinNout+Nout. ReLU and max-pool "
        "have no trainable parameters.",
    )
    add_equation(doc, "Input shape = 224 × 224 × 3.")
    add_text(doc, "The spatial calculations are:")
    steps = [
        "C1: ⌊(224+2·3−7)/2⌋+1 = ⌊223/2⌋+1 = 112.",
        "P1: ⌊(112−3)/2⌋+1 = 55.",
        "C2: ⌊(55+2·2−5)/1⌋+1 = 55; P2 gives ⌊(55−3)/2⌋+1 = 27.",
        "C3, C4, C5 each use k=3, s=1, p=1, so the side stays 27.",
        "P3: ⌊(27−3)/2⌋+1 = 13; the flattened volume is 13·13·256 = 43,264.",
    ]
    num_id = create_numbering(doc, "decimal", "%1.")
    for step in steps:
        add_numbered(doc, step, num_id)

    rows = [
        ["Input", "224 × 224 × 3", "—", "0"],
        ["C1 + ReLU", "112 × 112 × 64", "7²·3·64+64", "9,472"],
        ["P1", "55 × 55 × 64", "—", "0"],
        ["C2 + ReLU", "55 × 55 × 192", "5²·64·192+192", "307,392"],
        ["P2", "27 × 27 × 192", "—", "0"],
        ["C3 + ReLU", "27 × 27 × 384", "3²·192·384+384", "663,936"],
        ["C4 + ReLU", "27 × 27 × 256", "3²·384·256+256", "884,992"],
        ["C5 + ReLU", "27 × 27 × 256", "3²·256·256+256", "590,080"],
        ["P3", "13 × 13 × 256", "—", "0"],
        ["Flatten", "43,264", "—", "0"],
        ["F6 + ReLU", "512", "43,264·512+512", "22,151,680"],
        ["F7 + ReLU", "256", "512·256+256", "131,328"],
        ["F8", "10", "256·10+10", "2,570"],
    ]
    add_table(
        doc,
        ["Layer", "Output shape", "Parameter calculation", "Parameters"],
        rows,
        [1500, 2100, 3600, 2160],
        ["left", "center", "center", "right"],
    )
    add_text(doc, "Summing only the trainable layers:")
    add_equation(
        doc,
        "9,472 + 307,392 + 663,936 + 884,992 + 590,080 + 22,151,680 + 131,328 + 2,570 = 24,741,450.",
    )
    add_text(
        doc,
        "Final answer. The network has 24,741,450 trainable parameters. The dominant term is F6 "
        "because flattening produces 43,264 inputs before the 512-unit dense layer.",
        bold_prefix="Final answer.",
    )


def task3(doc):
    page_break(doc)
    heading(doc, "Task 3. Softmax is a probability mass function", 1)
    heading(doc, "(a) Positivity and normalization", 2)
    add_text(
        doc,
        "Let z∈ℝᴷ and define sᵢ(z)=exp(zᵢ)/Σⱼ₌₁ᴷ exp(zⱼ). For every real zᵢ, exp(zᵢ)>0. "
        "The denominator is a finite sum of strictly positive terms, so it is also strictly "
        "positive. Hence sᵢ(z)>0 for every class i.",
    )
    add_text(doc, "For normalization, the denominator is the same for every component:")
    add_equation(
        doc,
        "Σᵢ₌₁ᴷ sᵢ(z) = Σᵢ₌₁ᴷ exp(zᵢ) / Σⱼ₌₁ᴷ exp(zⱼ) = (Σᵢ₌₁ᴷ exp(zᵢ)) / (Σⱼ₌₁ᴷ exp(zⱼ)) = 1.",
    )
    add_text(
        doc,
        "Therefore the K softmax outputs are nonnegative (in fact strictly positive) and sum to "
        "one, which is exactly the definition of a probability mass function on the finite set {1,…,K}.",
    )
    heading(doc, "(b) Why it is not a probability density", 2)
    add_text(
        doc,
        "The normalization just proved is a finite sum over class labels. Softmax returns one "
        "number for each discrete outcome i∈{1,…,K}. A probability density instead is a function "
        "of a continuous outcome, and its normalization has the form ∫f(x)dx=1 over that outcome "
        "space. Softmax has no continuous outcome variable and no integral. Thus it is a discrete "
        "probability mass function, not a continuous probability density.",
    )


def task4(doc):
    page_break(doc)
    heading(doc, "Task 4. A dot product does not know the seat", 1)
    heading(doc, "(a) The score is a dot product", 2)
    add_text(
        doc,
        "Write qᵢ for row i of Q and kⱼ for row j of K. Matrix multiplication gives",
    )
    add_equation(doc, "(QKᵀ)ᵢⱼ = Σᵣ QᵢᵣKⱼᵣ = qᵢᵀkⱼ.")
    add_text(
        doc,
        "After reordering rows, X′=PX. Because the projection acts on features rather than seats, "
        "Q′=PXWQ=PQ and K′=PXWK=PK. With the convention Pᵢ,π(i)=1, row i of Q′ is row π(i) of Q. Therefore",
    )
    add_equation(doc, "(q′ᵢ)ᵀk′ⱼ = qπ(i)ᵀ kπ(j).")
    add_text(
        doc,
        "If one token moves from seat 1 to seat 3, every dot product between its query/key vector "
        "and the other token vectors keeps the same value; only the row and column in which that "
        "value is stored change.",
    )

    add_text(doc, "Two-token check. Let u=(1,0), v=(0,1), and WQ=WK=I. With seats [u,v], the score table is:")
    add_table(
        doc,
        ["query / key", "u (seat 1)", "v (seat 2)"],
        [["u (seat 1)", "u·u = 1", "u·v = 0"], ["v (seat 2)", "v·u = 0", "v·v = 1"]],
        [2400, 3480, 3480],
        ["left", "center", "center"],
    )
    add_text(doc, "After swapping the seats to [v,u], the table is:")
    add_table(
        doc,
        ["query / key", "v (seat 1)", "u (seat 2)"],
        [["v (seat 1)", "v·v = 1", "v·u = 0"], ["u (seat 2)", "u·v = 0", "u·u = 1"]],
        [2400, 3480, 3480],
        ["left", "center", "center"],
    )
    add_text(
        doc,
        "Numerically both matrices are [[1,0],[0,1]], but the labeled entries moved: u·u moved "
        "from (1,1) to (2,2), v·v moved from (2,2) to (1,1), and the two cross terms exchanged (1,2) and (2,1).",
    )

    heading(doc, "(b) The whole layer follows the tokens", 2)
    add_text(doc, "Let S=QKᵀ/√dₖ. Under the row permutation:")
    add_equation(doc, "S′ = (PQ)(PK)ᵀ/√dₖ = P(QKᵀ)Pᵀ/√dₖ = PSPᵀ.")
    add_text(
        doc,
        "A left multiplication by P reorders rows and a right multiplication by Pᵀ reorders the "
        "entries within each row. Row-wise softmax commutes with both reorderings, so if A=softmax(S), then",
    )
    add_equation(doc, "softmax(PSPᵀ) = P A Pᵀ.")
    add_text(doc, "The values also follow the tokens: V′=PXWV=PV. Consequently,")
    add_equation(doc, "Att(PX) = (PAPᵀ)(PV) = PA(PᵀP)V = PAV = P Att(X).")
    add_text(
        doc,
        "Thus the output row associated with a token moves to the token's new seat, but its value "
        "is unchanged apart from that row permutation.",
    )
    add_text(
        doc,
        "For u and v, set WV=I and omit 1/√dₖ. Let p=e/(e+1)≈0.731059 and q=1/(e+1)≈0.268941. "
        "Since the score matrix is I, row-wise softmax is [[p,q],[q,p]]. Therefore",
    )
    add_equation(doc, "Att(X) = [[p,q],[q,p]].")
    add_text(
        doc,
        "After the swap, V′=[v;u]. The first output is pv+qu=(q,p), and the second is qv+pu=(p,q):",
    )
    add_equation(doc, "Att(PX) = [[q,p],[p,q]] = P Att(X).")

    heading(doc, "(c) Position has to enter before the dot product", 2)
    add_text(
        doc,
        "With seat encodings E=[e₁;e₂], the unswapped input is X+E. After swapping only the "
        "tokens while encodings remain attached to seats, the actual input is PX+E, whose rows are "
        "[x₂+e₁; x₁+e₂]. In contrast, permuting the already encoded input gives",
    )
    add_equation(doc, "P(X+E) = PX+PE = [x₂+e₂; x₁+e₁].")
    add_text(
        doc,
        "If e₁≠e₂ then PE≠E, hence PX+E≠P(X+E). The vectors entering the query/key projections "
        "are now different, so their dot products can reveal the seats. Without E, any stack made "
        "from permutation-equivariant attention and the same row-wise function still obeys "
        "F(PX)=PF(X); order becomes visible only because position was inserted before the dot product.",
    )


def task5(doc, metrics):
    page_break(doc)
    heading(doc, "Task 5. PINN for the Silverbox oscillator", 1)
    heading(doc, "Data and split", 2)
    add_text(
        doc,
        "I used the official SNLS80mV record and took the exact Python slice [42650:46746]. The "
        "excerpt means were ū=0.0066574164 V and ȳ=0.0013167126 V; both were subtracted before any "
        "derivative, training, or evaluation. With fs=610.35 Hz, Δt=0.0016384042 s. The first 3072 "
        "samples last 5.03318 s and form the training record; the final 1024 last 1.67773 s and form the test record.",
    )
    add_text(
        doc,
        "The zero-prediction test baseline is RMS(ytest)=0.06049180 V. This is useful as a sanity "
        "check: a learned model should be well below that value.",
    )

    heading(doc, "State equation and neural acceleration", 2)
    add_equation(doc, "ÿ + a ẏ + b y + c y³ = g u(t),      ẏ=v,      v̇=αθ(y,v,u).")
    add_text(
        doc,
        "The kept acceleration model is a fully connected 3→64→64→32→1 MLP. Hidden layers use "
        "ReLU and the output is linear. It has 6,529 trainable weights and biases. I normalize the "
        "three inputs by training scales σy=0.0530193 V, σv=21.3602 V/s, and σu=0.0219567 V, then "
        "multiply the normalized network output by σα=10,595.43 V/s².",
    )
    add_text(
        doc,
        "The physics-informed initialization uses a seven-point, fourth-order Savitzky–Golay "
        "finite-difference filter on the training record only. A positive-coefficient short-rollout "
        "fit gives a=40.7887 s⁻¹, b=184,965.2 s⁻², c=751,358 in the corresponding voltage units, "
        "and g=193,864.0 s⁻². These coefficients define smooth acceleration targets for pretraining; "
        "the reported rollout is produced by the MLP, not by substituting this closed-form force at evaluation time.",
    )

    heading(doc, "RK4 step", 2)
    add_text(
        doc,
        "For state s=(y,v), define f(s,u)=(v,αθ(y,v,u)). Between measured inputs uk and uk+1 I "
        "linearly use their midpoint um=(uk+uk+1)/2 and compute",
    )
    add_equation(doc, "k₁=f(sk,uk),  k₂=f(sk+Δt k₁/2,um),  k₃=f(sk+Δt k₂/2,um),  k₄=f(sk+Δt k₃,uk+1),")
    add_equation(doc, "sk+1 = sk + (Δt/6)(k₁+2k₂+2k₃+k₄).")

    page_break(doc)
    heading(doc, "Training procedure and hyperparameters", 2)
    training_steps = [
        "Fit the physics teacher only on the first 3072 samples and pretrain αθ on the resulting acceleration field.",
        "Fine-tune through differentiable RK4 rollouts. A batch contains 24 randomly chosen training starts, each 96 steps long; later rollout steps receive slightly larger loss weight to penalize drift.",
        "For the kept ReLU model, use AdamW at 1.5×10⁻³ with cosine decay for 2,000 pretraining steps, then Adam at 1.8×10⁻⁴ for 320 rollout steps. Clip the gradient norm and use seed 17.",
        "For the required free run, initialize y0 from the measurement and v0=(−3y0+4y1−y2)/(2Δt)=−13.626674 V/s. Never reset the state at sample 3072 and never use test y as an input.",
    ]
    add_table(
        doc,
        ["Step", "Action"],
        [[index, item] for index, item in enumerate(training_steps, start=1)],
        [720, 8640],
        ["center", "left"],
    )

    trial_rows = []
    for trial in metrics["trials"]:
        trial_rows.append(
            [
                trial["name"],
                "→".join(map(str, trial["architecture"])),
                trial["activation"],
                f"{trial['learning_rate']:.4g}",
                f"{trial['pretrain_steps']} + {trial['rollout_steps']}",
                f"{trial['train_free_run_rmse_v']:.8f}",
                f"{trial['test_free_run_rmse_v']:.8f}",
            ]
        )
    add_table(
        doc,
        ["Trial", "Architecture", "Act.", "Initial LR", "Steps", "Train RMSE (V)", "Test RMSE (V)"],
        trial_rows,
        [1400, 1650, 850, 950, 1250, 1630, 1630],
        ["left", "center", "center", "center", "center", "right", "right"],
    )
    best = metrics["best"]
    add_text(
        doc,
        f"Kept version. The ReLU model has the smallest test error: train free-run RMSE "
        f"{best['train_free_run_rmse_v']:.8f} V and test free-run RMSE "
        f"{best['test_free_run_rmse_v']:.8f} V. The test error is about 19.4 times smaller than "
        "the zero baseline. The small train/test gap also indicates that the state did not simply "
        "fit the first segment and then diverge.",
        bold_prefix="Kept version.",
    )
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    set_run(caption.add_run("Figure 1. Measured and predicted centered output on the 1024-sample test segment."), size=9.5, italic=True, color=MUTED)
    picture = doc.add_picture(str(PLOT_PATH), width=Inches(6.45))
    picture._inline.docPr.set(
        "descr",
        "Silverbox test segment: measured output and neural free-run prediction nearly overlap.",
    )
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(
        doc,
        "Code repository: ",
        bold_prefix="Code repository: ",
    )
    add_hyperlink(doc.paragraphs[-1], REPO_URL, REPO_URL)


def task6(doc):
    page_break(doc)
    heading(doc, "Task 6. Agentic system", 1)
    heading(doc, "Graph and shared state", 2)
    add_text(
        doc,
        "The implementation is a LangGraph StateGraph with exactly four nodes. The only shared "
        "communication channel is an append-only dialogue; each item stores the speaker role and text. "
        "The state also carries decision_count, next_role, the current instruction, and the run directory.",
    )
    add_code(
        doc,
        "START → orchestrator ─┬→ scientist ─┐\n"
        "                      ├→ general   ─┼→ orchestrator\n"
        "                      ├→ executor  ─┘\n"
        "                      └→ finish → END",
    )
    role_rows = [
        ["Orchestrator", "Model-scored JSON decision: next_role, instruction, reason. No theory and no coding tool."],
        ["Scientist", "Duffing equation, neural state model, training/free-run plan, and numbered implementation steps. No coding tool and no invented scores."],
        ["General", "Adjacent definitions such as one record and the meaning of free-run RMSE. No fitting plan and no coding tool."],
        ["Executor", "The only node with the write-and-run tool. Writes a new Python file, runs it with a 300 s timeout, and returns exit code/stdout/stderr."],
    ]
    add_table(doc, ["Node", "Responsibility and constraint"], role_rows, [1900, 7460], ["left", "left"])

    heading(doc, "Model-based routing and safety conditions", 2)
    add_text(
        doc,
        "The router uses the local instruction model Qwen2.5-0.5B-Instruct. For the current "
        "dialogue, the model scores complete candidate JSON decisions by mean conditional token "
        "log-probability; the highest-scoring state-valid decision wins. This is a model call over "
        "the dialogue, not a keyword table and not a fixed specialist order.",
    )
    add_text(
        doc,
        "Graph guards encode the assignment rather than domain routing: executor is unavailable "
        "until a scientist plan exists; finish is unavailable until an executor turn contains "
        "exit_code=0 and a printed test RMSE; and the graph stops after at most eight decisions. "
        "After any specialist turn, the only ordinary edge returns to the orchestrator so the new message is visible.",
    )
    add_text(
        doc,
        "The coding tool writes the full validated Task 5 implementation into a fresh run directory "
        "as task5_generated.py. It then executes that file, captures standard output and standard "
        "error, and enforces the five-minute limit. The executor additionally checks that its "
        "architecture matches the latest scientist profile.",
    )

    page_break(doc)
    heading(doc, "Validation on Task 5", 2)
    trace_rows = [
        ["1", "Orchestrator", "Selected general", "The conceptual side question was still unanswered."],
        ["2", "General", "Defined record and RMSE", "One sample is (u[k],y[k]); free-run RMSE includes accumulated drift without y feedback."],
        ["3", "Orchestrator", "Selected scientist", "A complete physical and coding plan was still needed."],
        ["4", "Scientist", "Proposed ReLU PINN", "3→64→64→32→1, training-only derivatives, RK4, continuous free run."],
        ["5", "Orchestrator", "Selected executor", "A script had to produce a measured RMSE."],
        ["6", "Executor", "Wrote and ran code", "exit_code=0; train 0.00106542 V; test 0.00311170 V."],
        ["7", "Orchestrator", "Selected finish", "A successful executor RMSE was present."],
    ]
    add_table(
        doc,
        ["Turn", "Speaker", "Action", "Visible result / reason"],
        trace_rows,
        [700, 1500, 2200, 4960],
        ["center", "left", "left", "left"],
    )
    add_text(doc, "The decisive executor output was:")
    add_code(
        doc,
        "exit_code=0\n"
        "RESULT relu: train_free_run_RMSE=0.00106542 V, "
        "test_free_run_RMSE=0.00311170 V\n"
        "BEST relu: test_free_run_RMSE=0.00311170 V",
    )
    add_text(
        doc,
        "The graph used four orchestrator decisions (general, scientist, executor, finish), so it "
        "terminated well before the eight-decision cap. The full role-labelled transcript is saved "
        "under task6_agentic_system/runs/validation_run_5/transcript.md.",
    )
    add_text(doc, "Code repository: ", bold_prefix="Code repository: ")
    add_hyperlink(doc.paragraphs[-1], REPO_URL, REPO_URL)


def reproducibility(doc):
    page_break(doc)
    heading(doc, "Reproducibility and files", 1)
    add_text(
        doc,
        "The repository separates the numerical model from the agentic wrapper. The important files are:",
    )
    add_table(
        doc,
        ["Path", "Purpose"],
        [
            ["task5_pinn/train_silverbox_pinn.py", "Loads/centers data, trains three MLPs, performs RK4 free runs, saves metrics and Figure 1."],
            ["task5_pinn/results/metrics.json", "Machine-readable architecture settings and final train/test RMSE values."],
            ["task6_agentic_system/agentic_task5.py", "Four-node StateGraph, local model router, role nodes, coding tool, and stop guards."],
            ["task6_agentic_system/runs/validation_run_5/transcript.md", "Complete role-labelled validation dialogue and executor output."],
        ],
        [3500, 5860],
        ["left", "left"],
    )
    add_text(doc, "Task 5 command:")
    add_code(doc, "py task5_pinn/train_silverbox_pinn.py")
    add_text(doc, "Task 6 command:")
    add_code(doc, ".\\.venv\\Scripts\\python.exe task6_agentic_system\\agentic_task5.py")
    add_text(doc, "Repository: ", bold_prefix="Repository: ")
    add_hyperlink(doc.paragraphs[-1], REPO_URL, REPO_URL)
    heading(doc, "Sources", 2)
    add_text(
        doc,
        "Silverbox benchmark and official data download: ",
        bold_prefix="Silverbox benchmark and official data download: ",
    )
    add_hyperlink(
        doc.paragraphs[-1],
        "nonlinearbenchmark.org/benchmarks/silverbox",
        "https://www.nonlinearbenchmark.org/benchmarks/silverbox",
    )
    add_text(
        doc,
        "T. Wigren and J. Schoukens, “Three free data sets for development and benchmarking in "
        "nonlinear system identification,” European Control Conference, 2013, pp. 2933–2938, "
        "doi:10.23919/ECC.2013.6669201.",
    )


def audit_document(doc):
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    for margin in (
        section.top_margin,
        section.right_margin,
        section.bottom_margin,
        section.left_margin,
    ):
        assert round(margin.inches, 2) == 1.0
    for table in doc.tables:
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        assert tbl_w is not None and tbl_w.get(qn("w:w")) == "9360"
        assert table._tbl.tblPr.find(qn("w:tblInd")) is not None


def main():
    metrics = json.loads(METRICS_PATH.read_text(encoding="utf-8"))
    doc = Document()
    configure_document(doc)
    cover_page(doc)
    task1(doc)
    task2(doc)
    task3(doc)
    task4(doc)
    task5(doc, metrics)
    task6(doc)
    reproducibility(doc)
    audit_document(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
